"""
Document Ingestion Module for Smart Document Q&A Agent.

This module provides specialized reader functions for PDF, DOCX, and TXT files.
It handles file byte streams (from Streamlit uploads) and standard file paths,
extracting raw text along with metadata (source filename, page number).

Key concepts for project defense:
- PyMuPDF (fitz): High-performance C-based PDF rendering library used here for
  page-by-page text stream extraction.
- python-docx: Parses the OpenXML DOM tree of Microsoft Word (.docx) documents,
  extracting text across paragraphs and embedded data tables.
- Encoding resilience: Text file loading attempts UTF-8 first, with Latin-1
  fallback to prevent crashes on non-standard encodings.
"""

import io
import os
from typing import Any, Dict, List, Union


class DocumentParsingError(Exception):
    """Custom exception raised when document parsing encounters an unrecoverable error."""
    pass


def load_pdf(file_source: Union[str, bytes, io.BytesIO], filename: str) -> List[Dict[str, Any]]:
    """
    Extracts text page-by-page from a PDF document using PyMuPDF (fitz).

    Args:
        file_source: File path (str), raw bytes (bytes), or in-memory stream (BytesIO).
        filename: Original file name used for citation tracking.

    Returns:
        A list of dictionaries, where each entry represents one page:
        [{"text": str, "source": str, "page": int}, ...]

    Raises:
        DocumentParsingError: If PyMuPDF fails to open or parse the file.
    """
    import fitz  # PyMuPDF

    documents: List[Dict[str, Any]] = []

    try:
        # Open PDF from either raw bytes/stream or file path
        if isinstance(file_source, bytes):
            doc = fitz.open(stream=file_source, filetype="pdf")
        elif isinstance(file_source, io.BytesIO):
            doc = fitz.open(stream=file_source.getvalue(), filetype="pdf")
        elif hasattr(file_source, "read"):
            doc = fitz.open(stream=file_source.read(), filetype="pdf")
        elif isinstance(file_source, str):
            if not os.path.exists(file_source):
                raise FileNotFoundError(f"File not found at: {file_source}")
            doc = fitz.open(file_source)
        else:
            raise ValueError(f"Unsupported file source type: {type(file_source)}")

        # Extract text page by page (1-indexed for human readability)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text").strip()

            # Only record pages that contain non-whitespace text
            if text:
                documents.append({
                    "text": text,
                    "source": filename,
                    "page": page_num + 1
                })

        doc.close()

    except Exception as exc:
        raise DocumentParsingError(f"Failed to parse PDF '{filename}': {str(exc)}") from exc

    if not documents:
        raise DocumentParsingError(f"PDF '{filename}' contained no extractable text. (Might be a scanned image-only PDF).")

    return documents


def load_docx(file_source: Union[str, bytes, io.BytesIO], filename: str) -> List[Dict[str, Any]]:
    """
    Extracts text from a Microsoft Word (.docx) document using python-docx.

    Note on DOCX pagination:
    Unlike PDFs, Word documents are reflowable XML streams without fixed pages.
    Therefore, DOCX documents are treated as page 1 (or single logical units).

    Args:
        file_source: File path (str), raw bytes (bytes), or in-memory stream (BytesIO).
        filename: Original file name used for citation tracking.

    Returns:
        List containing a single dictionary with the full extracted document text.

    Raises:
        DocumentParsingError: If python-docx fails to read the document.
    """
    import docx

    try:
        # Normalize file stream for python-docx
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        elif isinstance(file_source, io.BytesIO):
            stream = file_source
        elif hasattr(file_source, "read"):
            stream = io.BytesIO(file_source.read())
        elif isinstance(file_source, str):
            stream = file_source
        else:
            raise ValueError(f"Unsupported file source type: {type(file_source)}")

        doc = docx.Document(stream)
        text_parts: List[str] = []

        # Extract regular paragraph text
        for paragraph in doc.paragraphs:
            clean_para = paragraph.text.strip()
            if clean_para:
                text_parts.append(clean_para)

        # Extract tabular text if present
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n\n".join(text_parts).strip()

        if not full_text:
            raise DocumentParsingError(f"Word document '{filename}' is empty or has no readable text.")

        return [{
            "text": full_text,
            "source": filename,
            "page": 1
        }]

    except Exception as exc:
        raise DocumentParsingError(f"Failed to parse DOCX '{filename}': {str(exc)}") from exc


def load_txt(file_source: Union[str, bytes, io.BytesIO], filename: str) -> List[Dict[str, Any]]:
    """
    Reads plain text (.txt) files with encoding resilience.

    Args:
        file_source: File path (str), raw bytes (bytes), or in-memory stream (BytesIO).
        filename: Original file name used for citation tracking.

    Returns:
        List containing a single dictionary with the document text.

    Raises:
        DocumentParsingError: If decoding fails or the file is empty.
    """
    try:
        raw_bytes: bytes

        if isinstance(file_source, bytes):
            raw_bytes = file_source
        elif isinstance(file_source, io.BytesIO):
            raw_bytes = file_source.getvalue()
        elif hasattr(file_source, "read"):
            data = file_source.read()
            raw_bytes = data if isinstance(data, bytes) else data.encode("utf-8")
        elif isinstance(file_source, str):
            with open(file_source, "rb") as f:
                raw_bytes = f.read()
        else:
            raise ValueError(f"Unsupported file source type: {type(file_source)}")

        # Try UTF-8 decoding first, then fallback to Latin-1
        try:
            text = raw_bytes.decode("utf-8").strip()
        except UnicodeDecodeError:
            text = raw_bytes.decode("latin-1", errors="replace").strip()

        if not text:
            raise DocumentParsingError(f"Text file '{filename}' is empty.")

        return [{
            "text": text,
            "source": filename,
            "page": 1
        }]

    except Exception as exc:
        raise DocumentParsingError(f"Failed to read TXT file '{filename}': {str(exc)}") from exc


def load_document(file_source: Any, filename: str) -> List[Dict[str, Any]]:
    """
    Unified entry point for loading documents by file extension.

    Supported formats:
    - .pdf  -> load_pdf
    - .docx -> load_docx
    - .txt  -> load_txt

    Args:
        file_source: File stream or path.
        filename: Name of the file with extension.

    Returns:
        List of parsed document page/unit dictionaries.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return load_pdf(file_source, filename)
    elif ext == ".docx":
        return load_docx(file_source, filename)
    elif ext == ".txt":
        return load_txt(file_source, filename)
    else:
        raise DocumentParsingError(
            f"Unsupported file format '{ext}' for file '{filename}'. Supported formats are: .pdf, .docx, .txt"
        )
