"""
Automated Test Suite for Smart Document Q&A Agent Pipeline.

This script rigorously validates:
1. Ingestion: TXT, PDF, and DOCX parsers.
2. Chunking: RecursiveCharacterTextSplitter (512 / 64) with metadata tracking.
3. Embeddings: SentenceTransformer all-MiniLM-L6-v2 (384 dimensions, L2-normalized).
4. Vector Store: FAISS IndexFlatL2 indexing and cosine similarity calculation.
5. Prompt Construction: Grounded prompt formatting with citations.
"""

import os
import sys
import io
import numpy as np

# Ensure project modules are importable
current_dir = os.path.dirname(os.path.abspath(__file__))
if current_dir not in sys.path:
    sys.path.insert(0, current_dir)

from ingestion import load_document, load_txt, load_pdf, load_docx, DocumentParsingError
from chunking import chunk_documents
from embeddings import get_embedding_model
from vector_store import FaissVectorStore
from qa_chain import format_retrieval_context, build_grounded_prompt


def test_txt_ingestion():
    print("[1/6] Testing TXT Ingestion...")
    sample_file = os.path.join(current_dir, "sample_docs", "artificial_intelligence_policy.txt")
    assert os.path.exists(sample_file), f"Sample file missing: {sample_file}"

    docs = load_document(sample_file, "artificial_intelligence_policy.txt")
    assert len(docs) == 1, f"Expected 1 document unit, got {len(docs)}"
    assert docs[0]["source"] == "artificial_intelligence_policy.txt"
    assert docs[0]["page"] == 1
    assert "academic integrity policy" in docs[0]["text"].lower()
    print("  -> TXT ingestion passed successfully.")
    return docs


def test_docx_ingestion():
    print("[2/6] Testing DOCX Generation and Ingestion...")
    import docx
    doc = docx.Document()
    doc.add_heading("Quantum Cryptography Briefing", level=1)
    doc.add_paragraph("Quantum Key Distribution (QKD) leverages the no-cloning theorem to guarantee secure key exchange.")
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Protocol"
    hdr_cells[1].text = "Year"
    row_cells = table.add_row().cells
    row_cells[0].text = "BB84"
    row_cells[1].text = "1984"

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    docs = load_docx(stream, "quantum_briefing.docx")
    assert len(docs) == 1
    assert "BB84" in docs[0]["text"]
    assert "no-cloning theorem" in docs[0]["text"]
    print("  -> DOCX ingestion passed successfully.")
    return docs


def test_pdf_ingestion():
    print("[3/6] Testing PDF Generation and Ingestion...")
    import fitz
    pdf_doc = fitz.open()
    # Create page 1
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 72), "Apex Research Lab - Page 1\nHigh performance computing overview.")
    # Create page 2
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 72), "Apex Research Lab - Page 2\nDistributed storage architectures and benchmarks.")

    pdf_bytes = pdf_doc.write()
    pdf_doc.close()

    docs = load_pdf(pdf_bytes, "research_lab_report.pdf")
    assert len(docs) == 2, f"Expected 2 pages, got {len(docs)}"
    assert docs[0]["page"] == 1
    assert "Page 1" in docs[0]["text"]
    assert docs[1]["page"] == 2
    assert "Page 2" in docs[1]["text"]
    print("  -> PDF ingestion passed successfully.")
    return docs


def test_chunking():
    print("[4/6] Testing Recursive Character Chunking...")
    sample_file = os.path.join(current_dir, "sample_docs", "quantum_computing_primer.txt")
    docs = load_document(sample_file, "quantum_computing_primer.txt")
    chunks = chunk_documents(docs, chunk_size=512, chunk_overlap=64)

    assert len(chunks) > 1, f"Expected multiple chunks, got {len(chunks)}"
    for chunk in chunks:
        assert "chunk_id" in chunk
        assert "text" in chunk
        assert "source" in chunk
        assert "page" in chunk
        assert len(chunk["text"]) <= 600, f"Chunk exceeded size threshold: {len(chunk['text'])}"

    print(f"  -> Chunking passed successfully: Created {len(chunks)} chunks.")
    return chunks


def test_embeddings_and_vector_store():
    print("[5/6] Testing Embeddings and FAISS Vector Store...")
    embedder = get_embedding_model("all-MiniLM-L6-v2")
    assert embedder.dimension == 384

    # Load and chunk all sample documents
    sample_dir = os.path.join(current_dir, "sample_docs")
    all_chunks = []
    for txt_file in os.listdir(sample_dir):
        if txt_file.endswith(".txt"):
            path = os.path.join(sample_dir, txt_file)
            docs = load_document(path, txt_file)
            all_chunks.extend(chunk_documents(docs, chunk_size=512, chunk_overlap=64))

    print(f"  -> Total chunks generated across sample docs: {len(all_chunks)}")
    texts = [c["text"] for c in all_chunks]
    embeddings = embedder.embed_texts(texts)

    assert embeddings.shape == (len(all_chunks), 384)
    # Check L2 normalization: norm should be approximately 1.0
    norms = np.linalg.norm(embeddings, axis=1)
    np.testing.assert_allclose(norms, 1.0, atol=1e-5)

    # Initialize FAISS Vector Store
    vector_store = FaissVectorStore(dimension=384)
    indexed_count = vector_store.add_documents(all_chunks, embeddings)
    assert indexed_count == len(all_chunks)

    stats = vector_store.get_stats()
    assert stats["total_chunks"] == len(all_chunks)
    assert stats["total_documents"] == 3

    # Query 1: Library undergraduate borrowing limit
    q1 = "What is the borrowing limit and loan period for undergraduate students?"
    q1_embed = embedder.embed_query(q1)
    results1 = vector_store.similarity_search(q1_embed, top_k=5)
    assert len(results1) == 5
    top_chunk1 = results1[0]
    print(f"  -> Top chunk source: {top_chunk1['source']}, similarity: {top_chunk1['similarity']}, text: {top_chunk1['text'][:80]}...")
    assert top_chunk1["source"] == "campus_library_guide.txt"
    assert top_chunk1["similarity"] > 0.40
    print(f"  -> Query 1 matched source '{top_chunk1['source']}' with cosine similarity: {top_chunk1['similarity']}")

    # Query 2: Shor's algorithm
    q2 = "Who developed Shor's algorithm and what does it solve?"
    q2_embed = embedder.embed_query(q2)
    results2 = vector_store.similarity_search(q2_embed, top_k=5)
    top_chunk2 = results2[0]
    assert top_chunk2["source"] == "quantum_computing_primer.txt"
    assert "Peter Shor" in top_chunk2["text"]
    assert top_chunk2["similarity"] > 0.40
    print(f"  -> Query 2 matched source '{top_chunk2['source']}' with cosine similarity: {top_chunk2['similarity']}")

    return vector_store, embedder, results1


def test_prompt_assembly(results1):
    print("[6/6] Testing Grounded Prompt Assembly...")
    context_str = format_retrieval_context(results1)
    assert "[Source 1: campus_library_guide.txt" in context_str

    prompt = build_grounded_prompt(
        question="What is the borrowing limit for undergraduates?",
        context=context_str
    )
    assert "Strict Operational Guidelines:" in prompt
    assert "I don't have enough information in the provided documents." in prompt
    assert "campus_library_guide.txt" in prompt
    print("  -> Prompt assembly passed successfully.")


def run_all_tests():
    print("==================================================")
    print("  RUNNING SMART DOCUMENT Q&A AGENT PIPELINE TESTS ")
    print("==================================================")
    test_txt_ingestion()
    test_docx_ingestion()
    test_pdf_ingestion()
    test_chunking()
    vector_store, embedder, results1 = test_embeddings_and_vector_store()
    test_prompt_assembly(results1)
    print("==================================================")
    print("  ALL 6 TEST SUITES COMPLETED SUCCESSFULLY!       ")
    print("==================================================")


if __name__ == "__main__":
    run_all_tests()
